"""Build evidence-based interim briefing and a research report working draft."""
from pathlib import Path
import json
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/"outputs/study_2026-10-23"
S=json.loads((ROOT/"data/processed/summary.json").read_text(encoding="utf-8"))
D=json.loads((ROOT/"data/processed/display_data.json").read_text(encoding="utf-8"))
REF=json.loads((ROOT/"docs/references.json").read_text(encoding="utf-8"))["sources"]
OUT.mkdir(parents=True,exist_ok=True)


def doc(title,subtitle):
    d=Document();sec=d.sections[0]
    sec.page_width=Inches(8.5);sec.page_height=Inches(11)
    sec.top_margin=sec.bottom_margin=Inches(.72);sec.left_margin=sec.right_margin=Inches(.78)
    for name,size in [("Normal",10.5),("Title",24),("Subtitle",11),("Heading 1",15),("Heading 2",12)]:
        st=d.styles[name];st.font.name="Aptos";st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0)
        st.paragraph_format.space_after=Pt(7)
        if name.startswith("Heading"):st.paragraph_format.space_before=Pt(13)
    d.styles["Normal"].paragraph_format.line_spacing=1.08
    for border in list(d.styles.element.xpath('.//w:pBdr')):
        border.getparent().remove(border)
    d.add_paragraph(title,"Title");d.add_paragraph(subtitle,"Subtitle")
    p=sec.footer.paragraphs[0];p.alignment=2
    p.add_run("COMP1110  |  23 October 2026  |  ")
    f=OxmlElement("w:fldSimple");f.set(qn("w:instr"),"PAGE");p._p.append(f)
    for r in p.runs:r.font.size=Pt(8);r.font.color.rgb=RGBColor.from_string("546575")
    return d


def p(d,text,style=None):return d.add_paragraph(text,style)
def h(d,text):return p(d,text,"Heading 1")
def table(d,heads,rows,widths=None,keep=False):
    t=d.add_table(rows=1,cols=len(heads));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    for i,text in enumerate(heads):t.rows[0].cells[i].text=str(text)
    for row in rows:
        cells=t.add_row().cells
        for i,text in enumerate(row):cells[i].text=str(text)
    for ri,row in enumerate(t.rows):
        for ci,c in enumerate(row.cells):
            if widths:c.width=Inches(widths[ci])
            c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr=c._tc.get_or_add_tcPr();b=OxmlElement("w:tcBorders")
            for side in ["top","left","bottom","right"]:
                el=OxmlElement("w:"+side);el.set(qn("w:val"),"single");el.set(qn("w:sz"),"4");el.set(qn("w:color"),"D9D9D9");b.append(el)
            pr.append(b);m=OxmlElement("w:tcMar")
            for side in ["top","left","bottom","right"]:
                el=OxmlElement("w:"+side);el.set(qn("w:w"),"90");el.set(qn("w:type"),"dxa");m.append(el)
            pr.append(m)
            if ri==0 or ri%2==0:
                sh=OxmlElement("w:shd");sh.set(qn("w:fill"),"E9EFF4" if ri==0 else "F7F9FB");pr.append(sh)
            for par in c.paragraphs:
                par.paragraph_format.space_after=Pt(2);par.paragraph_format.line_spacing=1
                if keep:par.paragraph_format.keep_with_next=True
                for r in par.runs:r.font.size=Pt(9);r.bold=ri==0
    rep=OxmlElement("w:tblHeader");t.rows[0]._tr.get_or_add_trPr().append(rep)
    p(d,"")
    return t


def status_rows(scope):
    return [[src,*[S["link_status"][scope][src].get(k,0) for k in ["Automatic proposal","Review needed","No strong candidate","No candidate"]]] for src in ["OSM","Overture"]]


speech=[
"Our project studies how reliably the same Hong Kong restaurant branch can be identified across FEHD, OpenStreetMap and Overture Places, and how comparable information agrees, differs or is missing.",
"We have now acquired the FEHD English and Chinese licence files dated 23 October. We have selected a reproducible 20-branch pilot and a provisional 100-branch main sample across our four districts. The pilot contains five branches per district. We also located FEHD's official spatial distribution: 99 of our 100 sampled licence IDs have a coordinate counterpart. We record its update dates rather than assuming every coordinate is current.",
"We have extracted OpenStreetMap food-related candidates and the relevant geographic subset of Overture's 21 October release. These raw extracts cover different categories, so we do not compare their total sizes as restaurant coverage. The matching workflow combines name, address and location evidence, and keeps close alternatives visible.",
f"In the pilot, the current rules propose {S['link_status']['pilot']['OSM'].get('Automatic proposal',0)} OSM links and {S['link_status']['pilot']['Overture'].get('Automatic proposal',0)} Overture links out of 20 branches. These are unvalidated algorithm proposals. The other cases need review or have no strong candidate. We cannot yet report matching accuracy or conclude that a source lacks those restaurants.",
"One useful diagnostic was a false candidate caused by a shopping-mall suffix in a longer restaurant name. We tightened the name rule and added a regression check. Other cases involve bilingual names, chain branches and addresses written at different levels of detail.",
"We are behind the original independent-review timetable. Sixty candidate pairs are prepared, with calibration and held-out branches separated, but two group members still need to label them independently. We also need to complete the pilot searches for missed candidates. After that, we can revise the rules, validate the matches and calculate field agreement with explicit denominators.",
"We would appreciate your advice on whether matching reliability, source coverage and name-and-address agreement provide sufficient depth if phone, hours and status have little comparable coverage. We also welcome feedback on how to interpret different address representations without treating every textual difference as an error."
]


brief=doc("STA Interim Progress Briefing","Research evidence and speaking notes 23 October 2026")
p(brief,"Use the figures below as the current catch-up position. Data acquisition and processing took place on 23 October; no work is backdated to the original schedule. Independent reference labels are still pending.")
h(brief,"Suggested spoken update")
for t in speech:p(brief,t)
brief.add_page_break();h(brief,"Evidence available for the meeting")
table(brief,["Item","Verified current position"],[['FEHD sampling frame','17,271 licence records; 7,593 in the four target district groups'],['Sample','20 pilot branches and 100 provisional main branches; five and 25 per district'],['Official spatial linkage','99 of 100 sampled licence IDs found in FEHD CSDI; all 99 normalised premises addresses agree with XML'],['External extracts','7,679 OSM food-related objects; 153,821 Overture places of all categories in the search box'],['Saved candidates',f"{S['candidate_pair_n']:,} ranked candidate pairs retained for 100 branches and two sources"],['Validation material','60 independent-review pairs; 21 calibration and 39 held-out, split by branch'],['Completed human review','0 pairs; no precision, recall or F1 is reported']], [1.55,5.3])
h(brief,"Pilot matching status")
table(brief,['Source','Automatic proposal','Review needed','No strong candidate','No candidate'],status_rows('pilot'),[1.05,1.55,1.4,1.6,1.25])
p(brief,"Counts describe the rule's decisions, not validated source coverage. A failed proposal may reflect naming, candidate retrieval or threshold choices. It does not prove a missing source record.")
h(brief,"Cases worth showing")
for t in ["Shake Shack in Sha Tin provides a relatively clear name-and-location example across the sources. It still requires independent branch verification.","COMEBUYTEA illustrates why the same brand elsewhere in Hong Kong cannot establish a branch match: some name hits are several kilometres from the FEHD premises.","UNAGI YONDAIME KIKUKAWA K11 MUSEA exposed a mall-name false positive. A shorter location suffix now cannot obtain the high name score that the earlier rule allowed."]:p(brief,t)
brief.add_page_break();h(brief,"Questions and immediate next work")
for t in ["Ask whether the three-source coverage and linkage investigation remains sufficiently substantive when the shared field set is narrow.","Ask whether exact normalised address equality should remain a representation measure, with semantic address review as a secondary analysis.","Ask how the STA would handle compound FEHD premises such as several named hotel facilities under one licence.","Two members complete the blinded 60-pair forms independently. Preserve original decisions and use a third member to adjudicate disagreements.","Complete 40 source-search entries covering all 20 pilot branches and both external sources; document any candidates found outside the generated set.","Freeze the calibrated rule before opening held-out labels. Retain the original pilot snapshot and record changes.","Complete the main sample review and freeze the analysis dataset by the planned 31 October target. Confirm all official dates and final-report instructions on Moodle."]:p(brief,t,"List Bullet")
h(brief,"中文汇报提醒")
for t in ["现在可以说：已取得三类来源、选好样本、跑过匹配、准备好60对核对材料，并发现和修正了一类误配规则。", "现在不能说：20家已经人工验证完成、某平台只收录了几家、准确率是多少，或者哪一来源更正确。", "如果问为什么落后：说明独立人工核对尚未完成，展示现有数据和接下来的具体分工即可，不需要编造此前做过的工作。", "这次材料由自动采集和辅助分析产生。组员应亲自看过试点和方法，能够解释采样、分母和误配案例，再在meeting中使用。"]:p(brief,t)
brief.save(OUT/"STA_Meeting_Briefing_2026-10-23.docx")

report=doc("Hong Kong Restaurant Record Consistency","Interim research report and final report working draft 23 October 2026")
p(report,"Team members: Qu Yunhao, Fan Yikun, Zhu Junning, He Mengzhe and Ding Shuai")
h(report,"Research status")
p(report,"We have acquired source snapshots, drawn a reproducible sample and implemented an initial matching workflow. The report records the evidence available on 23 October. Independent human labels, pilot missed-candidate searches and STA feedback remain outstanding. The results below are provisional and cannot establish matching accuracy or source correctness.")
h(report,"1 Research question and scope")
p(report,"A restaurant branch can have several records in administrative, community and compiled place datasets. These records describe different concepts and may use different names, premises detail and categories. We ask how reliably the same licensed Hong Kong branch can be identified across FEHD, OpenStreetMap and Overture Places, and how often comparable observed fields agree, differ or remain unavailable.")
p(report,"The study covers 100 provisional FEHD branches, 25 each in Central and Western, Wan Chai, Yau Tsim Mong and Sha Tin. Twenty branches, five per district, form the pilot. Equal district quotas support a bounded comparison rather than a representative Hong Kong estimate. Ratings, review text, spending and sentiment are outside the submitted study scope.")
h(report,"2 Related work and contribution")
p(report,"Sun et al. (2023) (Source 1) review established POI matching methods. Zhao et al. (2022) (Source 2) combine name, address and spatial similarity using a documented weighting method. These support an explainable linkage workflow, but neither establishes suitable local weights for our sample. Our initial weights are heuristic and require pilot calibration.")
p(report,"Yeow et al. (2021) (Source 3) examine POI validation in Singapore, including completeness and positional and thematic quality. Klinkhardt et al. (2023) (Source 4) compare OSM POIs against large-scale reference observations in Germany. Both motivate separate measures for missingness, linking and field comparison. Our project lacks their external reference coverage and cannot make equivalent accuracy claims.")
p(report,"Liu et al. (2023) (Source 5) use Hong Kong restaurant POIs to study catering resilience. This supplies a close geographical setting with a different research outcome. Compared with Yeow et al., our source combination and licensed-branch frame differ; compared with Liu et al., our outcome concerns record linkage and consistency rather than spatial resilience.")
p(report,"FEHD (Source 6) supplies the local administrative frame and an official spatial distribution. OpenStreetMap and Overpass (Source 7) support food-place extraction with optional contributed attributes. Overture Places (Source 8) supplies release-based place records, schema and upstream provenance. These data products make the study feasible while limiting which fields can be compared. Licensing and data release time do not show when a restaurant last changed its business information.")
p(report,"The existing Overmatch system (Source 9) already links OSM amenities to Overture and shows tag differences. Overture also publishes a government-restaurant matching tutorial. A discrepancy interface or a standard POI matching algorithm is therefore not our novelty claim. In the sources reviewed, we did not locate an evaluation of the three sources together using a reproducibly selected Hong Kong FEHD branch sample with explicit linkage uncertainty and field denominators. This is a bounded gap in the material reviewed, not proof that no such study exists.")
h(report,"3 Data and sampling")
table(report,['Source','Observed extract','Meaning and limits'],[['FEHD XML','23 October 2026; 17,271 licences','Names, premises addresses, licence type and expiry; no ratings or comparable daily hours'],['FEHD CSDI','99 sampled licence counterparts','Same official source; exact licence join and separately dated coordinates'],['OpenStreetMap','7,679 food-related objects in search box','Optional tags; way/relation centres; search box extends beyond Hong Kong'],['Overture Places','2026-10-21.0; 153,821 places in box','All categories retained; not the same population as food-only OSM']],[1.2,2.1,3.55])
p(report,"The search box is 113.80 to 114.45 degrees east and 22.15 to 22.60 degrees north. Broad extraction allows candidate searches around district borders. Neither external extract count is a count of Hong Kong restaurants. Raw responses, extraction SQL, source dates, actual retrieval times and checksums are retained. Unnecessary public OSM editor identifiers were removed with the transformation logged.")
p(report,"The FEHD four-district frame contains 7,593 eligible name/address groups: 1,787 Central and Western, 1,775 Wan Chai, 3,121 Yau Tsim Mong and 910 Sha Tin. FEHD Yau Tsim code 61 and Mong Kok code 62 are pooled explicitly. Exact normalised name/address duplicates would be merged with their licence IDs preserved; none were found in this frame. Alias or suite ambiguities may remain.")
p(report,"Within each district we order records by SHA256 of a fixed seed, district, normalised name and English address. The first five form the pilot and the first 25 the provisional main sample. Sampling precedes inspection of matching outcomes. The sample retains branches with poor matching yield. All new acquisition occurred on 23 October rather than the earlier planned dates.")
h(report,"4 Matching and validation method")
p(report,"Original fields are preserved. NFKC normalisation, case folding, punctuation and whitespace handling prepare comparison strings. Address rules standardise common English abbreviations, and supplied bilingual name variants remain available. These operations do not translate every address or resolve suite-level equivalence.")
p(report,"Candidates are generated from the union of a 750 m geographic neighbourhood, exact name hits and limited fuzzy-name fallback searches. FEHD CSDI coordinates are joined by SEARCH02_EN to XML LICNO; all 99 returned records agree on the normalised English premises address. The points remain supporting evidence with their own update dates. No commercial geocoding is used.")
p(report,"The initial score weights name 0.55, address 0.25 and distance 0.20, renormalising when a feature is unavailable. Geographic similarity decreases with distance using a 150 m scale. FEHD has no phone field, so phone does not contribute to FEHD-to-source scores. These values are starting rules rather than probabilities or locally validated parameters.")
p(report,"An automatic proposal requires score at least 0.82, name similarity at least 0.84, distance no more than 200 m and at least a 0.08 margin over the next candidate. Duplicate assignments require review. A location suffix alone cannot obtain a high token-set name score when it is much shorter than the business name. The saved ranked table contains the top 12 candidates for each branch/source combination, with 2,206 rows retained in total.")
p(report,"The independent-review set contains 60 pairs: each pilot branch's top candidates and one competing candidate. Seven branches supply 21 calibration pairs; 13 supply 39 label-held-out pairs. The inspected K11 development case is reserved for calibration before human labelling. Other assignments follow a fixed branch hash. Pilot features were inspected, so this is a reference-label holdout, not an untouched external test. Scores and assignments are hidden from reviewers. This purposeful set is not prevalence-weighted.")
p(report,"Two members must independently choose Match, Non-match or Uncertain and record their actual identity, date and evidence. Disagreements require documented adjudication without overwriting the initial labels. Calibration precedes freezing the rule and opening held-out results. A separate 40-row search form covers both external sources for all 20 pilot branches to examine candidates missed by the procedure.")
p(report,"Precision, recall and F1 will be calculated only from an explicit labelled evaluation universe with uncertain cases counted separately. Pair-classification recall within that set is distinct from candidate-retrieval recall. Currently all three performance metrics remain unavailable because independent reference labels have not been supplied.")
h(report,"5 Provisional observations")
table(report,['Source','Automatic proposal','Review needed','No strong candidate','No candidate'],status_rows('pilot'),[1.05,1.55,1.4,1.6,1.25])
fig=OUT/'figures/pilot_link_status.png'
if fig.exists():report.add_picture(str(fig),width=Inches(6.75))
p(report,"The initial rules produce three OSM proposals and seven Overture proposals among the 20 pilot branches. These describe algorithmic decisions, not verified coverage. A weak result can arise from a missing alias, stale or incomplete attributes, a competing branch, the candidate procedure or the thresholds. It does not establish that the source lacks the branch.")
p(report,"For the 100-branch provisional sample, the same rules propose 12 OSM and 37 Overture links. OSM has 36 review-needed, 48 no-strong-candidate and four no-candidate cases. Overture has 40, 22 and one respectively. Review may change these counts, so the analysis dataset is not frozen.")
primary=[x for x in S['field_summary'] if x['scope']=='pilot' and x['field'] in ['name','address']]
table(report,['Pair','Field','Agree','Different text','Comparable n','Other cases'],[[x['pair'],x['field'],x['agree'],x['disagree'],x['comparable_n'],x['n']-x['comparable_n']] for x in primary],[1.65,.75,.7,1.15,1.1,1.5],keep=True)
p(report,"Only automatically proposed links enter these provisional field comparisons. Agreement uses the explicit comparable denominator; unavailable fields and uncertain links are excluded and reported. Exact normalised address equality is a representation measure. Different text often reflects administrative premises descriptions versus shorter place addresses, so it cannot support a claim that a location is wrong. Name and address evidence also select the links, introducing selection bias into their conditional agreement.")
h(report,"6 Diagnostic cases and design response")
p(report,"COMEBUYTEA demonstrates chain ambiguity: strong brand-name hits can refer to branches several kilometres away. The display therefore retains candidate distances and alternative IDs. A near point plus an uninformative name is also insufficient to establish branch identity.")
p(report,"The pilot UNAGI YONDAIME KIKUKAWA K11 MUSEA initially attracted the mall itself because token-set overlap rewarded its location suffix. Restricting that operation by relative name length removes the high-name-score route for this case. A regression check preserves the correction, while independent review is still needed to judge real matches.")
p(report,"Shake Shack in Sha Tin gives a comparatively clear name-and-location example. Several hotel and multi-venue licences show a different issue: an administrative licence may cover more than one customer-facing facility. Those cases need an explicit granularity decision rather than forced one-to-one links.")
p(report,"The interface displays original source values, current link status, candidate IDs, feature scores and retrieval context. Agree, Disagree, Unavailable, Non-comparable and Ambiguous match remain distinct. A separate blinded review page permits actual reviewers to save and export decisions; it does not silently turn browser entries into validated research results.")
h(report,"7 Limitations ethics and data use")
p(report,"This small equal-quota sample cannot establish Hong Kong-wide coverage or causal patterns. Official licences are an administrative frame, not a census of operating dining venues. CSDI and the external sources have different update cycles, positional meanings and category definitions. OSM's food-only extraction can miss differently tagged or disused objects; Overture's all-place extraction can surface non-restaurant candidates.")
p(report,"FEHD licence type is not cuisine. Overture operating_status is not a daily schedule, and existence confidence is not field accuracy. Overture Places does not provide a general opening-hours field; optional OSM hours do not create a common comparison by themselves. These fields remain secondary or non-comparable until suitable values and definitions exist.")
p(report,"The study collects business and administrative observations rather than customer reviews. Attribute FEHD and DATA.GOV.HK, OpenStreetMap contributors under ODbL, and Overture with its per-record upstream sources and applicable licences. Code and data licensing are separate. No source is described as universally correct or officially endorsing the project.")
h(report,"8 Remaining work and final submission")
p(report,"The immediate priorities are authentic independent labels, adjudication, pilot missed-candidate searches and STA feedback on scope and address interpretation. The team can then freeze the main dataset, calculate validated coverage and conditional field summaries, review sensitivity to matching rules and complete the final discussion. Dates should be checked on Moodle; the supplied general guideline lists 31 October as no course-wide deadline, but that is this group's planned dataset-freeze target.")
p(report,"Individual progress logs and self-reflections must describe each student's actual contributions and learning. They have not been manufactured from these automated outputs. Separate final-report and presentation rubrics were not supplied with the two initial course PDFs, so final formatting and submission checks must use those later instructions when available.")
h(report,"References")
for i,r in enumerate(REF,1):
    authors=", ".join(r.get('authors',[])) or r.get('organisation') or r.get('author') or r.get('title','')
    year=r.get('year') or 'n.d.'
    venue=r.get('venue','')
    publication=venue
    if r.get('volume'):publication+=', '+r['volume']+(('('+r['issue']+')') if r.get('issue') else '')
    if r.get('pages') or r.get('article_number'):publication+=', '+r.get('pages',r.get('article_number',''))
    text=f"Source {i}. {authors} ({year}). {r['title']}. "+(publication+'. ' if publication else '')+r['url']
    par=p(report,text)
    for run in par.runs:run.font.size=Pt(9)
    par.paragraph_format.space_after=Pt(6)
report.save(OUT/"Research_Report_Working_Draft.docx")
print("Created briefing and research report working draft",flush=True)
