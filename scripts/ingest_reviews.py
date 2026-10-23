"""Validate genuine reviewer files; expose calibration labels only, never invent labels."""
import argparse
import csv
from datetime import date, datetime, timezone
import hashlib
import json
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
DECISIONS = {"Match", "Non-match", "Uncertain"}


def text(value):
    return "" if value is None else str(value).strip()


def actual_date(value):
    if isinstance(value, datetime):
        result = value.date()
    elif isinstance(value, date):
        result = value
    else:
        try:
            result = datetime.fromisoformat(text(value).replace("Z", "+00:00")).date()
        except ValueError as exc:
            raise ValueError("Use an actual ISO review date, e.g. YYYY-MM-DD.") from exc
    if result > datetime.now().astimezone().date():
        raise ValueError("Review date cannot be in the future.")
    return result.isoformat()


def load_rows(path):
    if path.suffix.lower() == ".csv":
        with path.open(encoding="utf-8-sig", newline="") as stream:
            return list(csv.DictReader(stream))
    if path.suffix.lower() != ".xlsx":
        raise ValueError("Use a completed CSV or XLSX file.")
    from openpyxl import load_workbook
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = book["Review"]
        rows = list(sheet.iter_rows(min_row=8, max_row=67, values_only=True))
        return [{"pair_id": r[0], "candidate_id": r[6], "decision": r[8],
                 "evidence_and_reason": r[9], "reviewer_name": r[10], "reviewed_at": r[11]}
                for r in rows]
    finally:
        book.close()


def validate(rows, admin):
    by_id = {text(r.get("pair_id")): r for r in rows}
    if len(rows) != 60 or len(by_id) != 60 or set(by_id) != set(admin):
        raise ValueError("Exactly the 60 original, unique pair IDs are required.")
    result = {}
    for pair_id, row in by_id.items():
        if text(row.get("candidate_id")) != admin[pair_id]["candidate_id"]:
            raise ValueError(f"Candidate ID was changed for {pair_id}.")
        decision = text(row.get("decision"))
        name = text(row.get("reviewer_name"))
        reason = text(row.get("evidence_and_reason"))
        if decision not in DECISIONS or not name or not reason:
            raise ValueError(f"Incomplete/invalid decision, name or evidence for {pair_id}.")
        result[pair_id] = {"decision": decision, "reviewer_name": name,
                           "evidence_and_reason": reason, "reviewed_at": actual_date(row.get("reviewed_at"))}
    if len({r["reviewer_name"] for r in result.values()}) != 1:
        raise ValueError("Each independent file must identify one actual reviewer consistently.")
    return result


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("reviewer_a", type=Path)
    parser.add_argument("reviewer_b", type=Path)
    parser.add_argument("--output-dir", type=Path, required=True,
                        help="A new directory; original reviewer files are never overwritten.")
    args = parser.parse_args()
    admin = {r["pair_id"]: r for r in json.loads((ROOT / "data/review/review_admin.json").read_text(encoding="utf-8"))}
    a, b = [validate(load_rows(p), admin) for p in [args.reviewer_a, args.reviewer_b]]
    if next(iter(a.values()))["reviewer_name"].casefold() == next(iter(b.values()))["reviewer_name"].casefold():
        raise ValueError("Two different actual human reviewers are required.")
    rows = []
    for pair_id, original in admin.items():
        if original["split"] != "calibration":
            continue
        row = {k: original[k] for k in ["pair_id", "branch_id", "candidate_source", "candidate_id"]}
        for slot, labels in [("A", a), ("B", b)]:
            row.update({slot + "_" + k: v for k, v in labels[pair_id].items()})
        row.update({"needs_adjudication": a[pair_id]["decision"] != b[pair_id]["decision"] or "Uncertain" in [a[pair_id]["decision"], b[pair_id]["decision"]],
                    "adjudicator_name": "", "final_decision": "", "adjudication_evidence": "", "adjudicated_at": ""})
        rows.append(row)
    args.output_dir.mkdir(parents=True, exist_ok=False)
    with (args.output_dir / "calibration_adjudication.csv").open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=list(rows[0]))
        writer.writeheader(); writer.writerows(rows)
    meta = {"ingested_utc": datetime.now(timezone.utc).isoformat(timespec="seconds"),
            "inputs": [{"file": p.name, "sha256": hashlib.sha256(p.read_bytes()).hexdigest()} for p in [args.reviewer_a, args.reviewer_b]],
            "validated_rows_per_reviewer": 60, "calibration_rows_exported": len(rows),
            "held_out_labels_exported": 0,
            "next_step": "Adjudicate calibration, document any changes and freeze rules before opening held-out labels. This command does not verify independence or calculate accuracy."}
    (args.output_dir / "ingestion.json").write_text(json.dumps(meta, indent=2), encoding="utf-8")
    print("Validated two complete files. Exported calibration rows only; held-out labels and metrics remain unopened.")


if __name__ == "__main__":
    main()
